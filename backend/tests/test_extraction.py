import os
import tempfile
from pathlib import Path
import pytest
import docx

from app.extraction import (
    validate_file,
    clean_text,
    extract_from_txt,
    extract_from_docx,
    extract_document,
    ALLOWED_EXTENSIONS
)


def test_clean_text():
    raw = "  Hello   world \r\n\n\n\nThis is a test. \u200B "
    cleaned = clean_text(raw)
    assert cleaned == "Hello world\n\nThis is a test."


def test_validate_file_extensions():
    valid, _ = validate_file("resume.pdf")
    assert valid is True

    valid, _ = validate_file("document.docx")
    assert valid is True

    valid, _ = validate_file("notes.txt")
    assert valid is True

    valid, err = validate_file("virus.exe")
    assert valid is False
    assert "Unsupported file extension" in err


def test_validate_file_size():
    with tempfile.NamedTemporaryFile(suffix=".txt", delete=False) as f:
        f.write(b"Small file content")
        tmp_name = f.name

    try:
        valid, _ = validate_file(tmp_name, file_size=500)
        assert valid is True

        valid_large, err = validate_file(tmp_name, file_size=20 * 1024 * 1024)
        assert valid_large is False
        assert "exceeds maximum allowed size" in err
    finally:
        if os.path.exists(tmp_name):
            os.remove(tmp_name)


def test_extract_from_txt():
    with tempfile.NamedTemporaryFile(suffix=".txt", delete=False, mode="w", encoding="utf-8") as f:
        f.write("Line 1: Pavan Dasam\n\nLine 2: AI Engineer specializing in RAG.")
        tmp_name = f.name

    try:
        text, meta = extract_from_txt(tmp_name)
        assert "Pavan Dasam" in text
        assert "AI Engineer" in text
        assert meta["line_count"] >= 2
    finally:
        if os.path.exists(tmp_name):
            os.remove(tmp_name)


def test_extract_from_docx():
    doc = docx.Document()
    doc.add_heading("Pavan Dasam Profile", level=1)
    doc.add_paragraph("Agentic AI Engineer with LangGraph experience.")

    with tempfile.NamedTemporaryFile(suffix=".docx", delete=False) as f:
        tmp_name = f.name

    try:
        doc.save(tmp_name)
        text, meta = extract_from_docx(tmp_name)
        assert "Pavan Dasam Profile" in text
        assert "Agentic AI Engineer" in text
    finally:
        if os.path.exists(tmp_name):
            os.remove(tmp_name)
