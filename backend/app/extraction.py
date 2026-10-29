import os
import re
from pathlib import Path
from typing import Tuple, Dict, Any
import docx
import pymupdf


ALLOWED_EXTENSIONS = {".pdf", ".docx", ".txt"}
MAX_FILE_SIZE_BYTES = 15 * 1024 * 1024  # 15 MB


def validate_file(file_path: str or Path, file_size: int = None) -> Tuple[bool, str]:
    path = Path(file_path)
    ext = path.suffix.lower()

    if ext not in ALLOWED_EXTENSIONS:
        return False, f"Unsupported file extension '{ext}'. Allowed extensions are: {', '.join(sorted(ALLOWED_EXTENSIONS))}"

    actual_size = file_size if file_size is not None else (os.path.getsize(path) if path.exists() else 0)
    if actual_size > MAX_FILE_SIZE_BYTES:
        max_mb = MAX_FILE_SIZE_BYTES / (1024 * 1024)
        return False, f"File exceeds maximum allowed size of {max_mb:.0f} MB (got {actual_size / (1024 * 1024):.2f} MB)"

    return True, ""


def clean_text(text: str) -> str:
    """Normalize whitespace while preserving paragraphs and punctuation."""
    if not text:
        return ""
    # Normalize carriage returns
    text = text.replace("\r\n", "\n").replace("\r", "\n")
    # Replace weird unicode spaces
    text = re.sub(r"[\u200B-\u200D\uFEFF]", "", text)
    # Strip whitespace from each line
    lines = [re.sub(r"[ \t]+", " ", line).strip() for line in text.split("\n")]
    text = "\n".join(lines)
    # Collapse more than two consecutive newlines into two (paragraph break)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


def extract_from_pdf(file_path: str or Path) -> Tuple[str, Dict[str, Any]]:
    doc = pymupdf.open(str(file_path))
    pages_text = []
    metadata = {
        "page_count": len(doc),
        "title": doc.metadata.get("title") or "",
        "author": doc.metadata.get("author") or "",
    }
    for i, page in enumerate(doc):
        text = page.get_text("text") or ""
        cleaned = clean_text(text)
        if cleaned:
            pages_text.append(cleaned)
    doc.close()
    full_text = "\n\n".join(pages_text)
    return full_text, metadata


def extract_from_docx(file_path: str or Path) -> Tuple[str, Dict[str, Any]]:
    doc = docx.Document(str(file_path))
    paragraphs = []
    for p in doc.paragraphs:
        cleaned = clean_text(p.text)
        if cleaned:
            paragraphs.append(cleaned)

    # Also extract text from any tables in the DOCX
    for table in doc.tables:
        for row in table.rows:
            row_cells = [clean_text(cell.text) for cell in row.cells if clean_text(cell.text)]
            if row_cells:
                paragraphs.append(" | ".join(row_cells))

    full_text = "\n\n".join(paragraphs)
    metadata = {
        "paragraph_count": len(doc.paragraphs),
        "table_count": len(doc.tables)
    }
    return full_text, metadata


def extract_from_txt(file_path: str or Path) -> Tuple[str, Dict[str, Any]]:
    encodings = ["utf-8", "utf-8-sig", "latin-1", "cp1252"]
    full_text = ""
    used_encoding = "utf-8"
    for enc in encodings:
        try:
            with open(file_path, "r", encoding=enc) as f:
                full_text = f.read()
                used_encoding = enc
                break
        except UnicodeDecodeError:
            continue

    cleaned = clean_text(full_text)
    metadata = {
        "encoding": used_encoding,
        "line_count": len(cleaned.splitlines())
    }
    return cleaned, metadata


def extract_document(file_path: str or Path) -> Tuple[str, Dict[str, Any]]:
    path = Path(file_path)
    valid, err = validate_file(path)
    if not valid:
        raise ValueError(err)

    ext = path.suffix.lower()
    if ext == ".pdf":
        text, meta = extract_from_pdf(path)
    elif ext == ".docx":
        text, meta = extract_from_docx(path)
    elif ext == ".txt":
        text, meta = extract_from_txt(path)
    else:
        raise ValueError(f"Unsupported file extension: {ext}")

    meta["file_name"] = path.name
    meta["file_size"] = os.path.getsize(path)
    meta["extension"] = ext
    return text, meta
